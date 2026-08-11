import docx
from docx.shared import Inches
import struct, zlib

def make_png():
    sig = b'\x89PNG\r\n\x1a\n'
    def chunk(t, d):
        c = t + d
        return struct.pack('>I', len(d)) + c + struct.pack('>I', zlib.crc32(c) & 0xffffffff)
    ihdr = struct.pack('>IIBBBBB', 1, 1, 8, 2, 0, 0, 0)
    raw = b'\x00\xff\x00\x00'
    idat = zlib.compress(raw)
    return sig + chunk(b'IHDR', ihdr) + chunk(b'IDAT', idat) + chunk(b'IEND', b'')

png = make_png()
with open('build/test_img.png', 'wb') as f:
    f.write(png)

d = docx.Document()
d.add_paragraph('Paragraph 1: this is a test doc for dragging an image across pages. Drag the image below to page 2 and check for errors.')
for _ in range(6):
    d.add_paragraph('Fill text to push the image near the bottom of page 1 for cross-page drag testing. ' * 2)
d.add_picture('build/test_img.png', width=Inches(2))
d.add_paragraph('Paragraph 2: content after the image.')
for _ in range(6):
    d.add_paragraph('More fill text to ensure multiple pages exist for the drag test. ' * 2)
d.save('build/test_drag_image.docx')
print('saved build/test_drag_image.docx')
