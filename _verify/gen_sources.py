"""生成对比测试源文件：pptx / xlsx（现代格式），后续用 LibreOffice 转成 ppt / xls（老版二进制）。"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from openpyxl import Workbook

HERE = os.path.dirname(os.path.abspath(__file__))


def gen_pptx():
    prs = Presentation()
    # 第 1 页
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "第一页标题Alpha"
    tf = s1.placeholders[1].text_frame
    tf.text = "副标题Beta"
    # 第 2 页
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "第二页标题中文测试"
    tf2 = s2.placeholders[1].text_frame
    for i, t in enumerate(["要点一 One", "要点二 中文Two", "要点三 Three"]):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.text = t
    # 第 3 页：含混合英文+中文标题，验证编码
    s3 = prs.slides.add_slide(prs.slide_layouts[1])
    s3.shapes.title.text = "Wang五与混合ABC测试"
    tf3 = s3.placeholders[1].text_frame
    tf3.text = "Hello世界 123"
    out = os.path.join(HERE, "source.pptx")
    prs.save(out)
    print("wrote", out)


def gen_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "销售表"
    ws.append(["姓名", "部门", "销售额"])
    ws.append(["Wang五", "销售", 12000])
    ws.append(["李四", "市场", 9800])
    ws.append(["张三", "销售", 15000])
    # 第二个工作表
    ws2 = wb.create_sheet("汇总")
    ws2.append(["总计", 36800])
    ws2.append(["平均", 12266.67])
    out = os.path.join(HERE, "source.xlsx")
    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    import traceback
    try:
        gen_pptx()
        gen_xlsx()
    except Exception:
        with open(os.path.join(HERE, "gen_err.log"), "w", encoding="utf-8") as f:
            traceback.print_exc(file=f)
        raise
